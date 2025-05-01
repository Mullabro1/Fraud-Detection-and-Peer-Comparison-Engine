import os
import json
import openpyxl
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter

# Define paths
script_directory = os.path.dirname(os.path.abspath(__file__))
json_folder = os.path.join(script_directory, "json")
calc_file = os.path.join(json_folder, "calc2.json")
template_file = os.path.join(script_directory, "template", "ratio_temp.json")
sector_file = os.path.join(script_directory, "template", "sector.json")

# Output only Current_Ratio for now
excel_path = os.path.join(json_folder, "current_ratio.xlsx")

# Color map
COLOR_MAP = {
    "red": PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid"),
    "yellow": PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid"),
    "green": PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid"),
    "blue": PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid"),
    "gray": PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid"),
}

# Styles per column
column_styles = {
    1: {"alignment": Alignment(horizontal='center', vertical='center'), "font": Font(bold=True), "width": 25, "row_height": 25},
    2: {"alignment": Alignment(horizontal='right', vertical='center'), "font": Font(italic=True), "width": 13, "row_height": 25},
    3: {"alignment": Alignment(horizontal='left', vertical='center'), "font": Font(underline='single'), "width": 25, "row_height": 25},
    4: {"alignment": Alignment(horizontal='center', vertical='top'), "font": Font(bold=True, italic=True), "width": 55, "row_height": 25},
    5: {"alignment": Alignment(horizontal='center', vertical='center'), "font": Font(bold=True, color="FF0000"), "width": 11, "row_height": 25},
    6: {"alignment": Alignment(horizontal='center', vertical='bottom'), "font": Font(underline='single'), "width": 55, "row_height": 25},
    7: {"alignment": Alignment(horizontal='left', vertical='center'), "font": Font(bold=True, italic=True), "width": 55, "row_height": 25},
    8: {"alignment": Alignment(horizontal='center', vertical='center'), "font": Font(italic=True), "width": 18, "row_height": 25}
}

# Load JSON files
with open(calc_file, "r") as f:
    calc_data = json.load(f)

with open(template_file, "r") as f:
    templates = json.load(f)

with open(sector_file, "r") as f:
    sector_data = json.load(f)

# Extract sector and ratio
selected_sector = sector_data["sector"]
ratio_name = "Current_Ratio"

# Validate and extract template for selected sector
if selected_sector not in templates[ratio_name]:
    raise ValueError(f"Sector '{selected_sector}' not found in templates.")

sector_template = templates[ratio_name][selected_sector]

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = ratio_name

# Extract sector and ratio
selected_sector = sector_data["sector"]
ideal_range = sector_data["ideal_range"]  # Get ideal range from sector.json
ratio_name = "Current_Ratio"

# Header
ws.append([f"{ratio_name.replace('_', ' ')}"])
ws.append(["Ideal Range:\n", ideal_range])
ws.append([])  # Spacer
ws.append(["Thresholds for \n"])
ws.append([f"{selected_sector.replace('_', ' ')}"])
ws.append(["Industry"])
ws.append([])  # Spacer


# Interpretation Table
ws.append(["Range", "Status", "Short Description", "Interpretation"])
for i, t in enumerate(sector_template):
    threshold = float("inf") if t["threshold"] == "inf" else t["threshold"]
    prev = sector_template[i - 1]["threshold"] if i > 0 else None
    range_text = (
        f"≤ {threshold}" if i == 0 else
        f"> {prev}" if threshold == float("inf") else
        f"{prev} ≤ {threshold}"
    )
    row = [range_text, t["status"], t["short_description"], t["interpretation"]]
    ws.append(row)
    for col in range(1, 5):
        ws.cell(row=ws.max_row, column=col).fill = COLOR_MAP[t["color"]]

# Spacer before data
ws.append([])
ws.append(["Standalone"])
ws.append(["Org ID", "Data Range", "Company Value", "Short Description", "Status", "Interpretation"])

for entry in calc_data:
    for period in entry["periods"]:
        value = period.get(ratio_name)
        if value is not None:
            for t in sector_template:
                threshold = float("inf") if t["threshold"] == "inf" else t["threshold"]
                if value <= threshold:
                    row = [
                        entry["org_id"],
                        period["data_range"],
                        round(value, 2),
                        t["short_description"],  # moved before status
                        t["status"],
                        t["interpretation"]
                    ]
                    ws.append(row)
                    for col in range(1, 7):
                        ws.cell(row=ws.max_row, column=col).fill = COLOR_MAP[t["color"]]
                    break

# Spacer between tables
ws.append([])
ws.append(["Consolidated"])
ws.append(["Org ID", "Data Range", "Company Value", "Short Description", "Status", "Interpretation"])

for entry in calc_data:
    for period in entry["periods"]:
        value = period.get(ratio_name)
        if value is not None:
            for t in sector_template:
                threshold = float("inf") if t["threshold"] == "inf" else t["threshold"]
                if value <= threshold:
                    row = [
                        entry["org_id"],
                        period["data_range"],
                        round(value, 2),
                        t["short_description"],  # moved before status
                        t["status"],
                        t["interpretation"]
                    ]
                    ws.append(row)
                    for col in range(1, 7):
                        ws.cell(row=ws.max_row, column=col).fill = COLOR_MAP[t["color"]]
                    break

# Apply styles
for col_idx, style in column_styles.items():
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = style["width"]
    for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.alignment = style["alignment"]
            cell.font = style["font"]
    if "row_height" in style:
        ws.row_dimensions[1].height = style["row_height"]

# Save
wb.save(excel_path)
print("✅ Excel report generated using selected sector's thresholds.")
